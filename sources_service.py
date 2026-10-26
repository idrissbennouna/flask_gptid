"""Gestion des sources documentaires : extraction, découpage en passages et recherche BM25."""
import io
import math
import re
import unicodedata
from collections import Counter

import database as db

ALLOWED_EXTENSIONS = {".txt", ".md", ".csv", ".json", ".pdf", ".docx"}

STOPWORDS = set(
    "le la les un une des du de d l et ou en au aux ce ces cet cette dans par pour sur avec sans que qui "
    "est sont etre ete a il elle ils elles nous vous je tu on se sa son ses leur leurs mais donc or ni car "
    "the of and to in is are for on with that this it as by be at or an from".split()
)


# ----------------------------------------------------------------------------
# Extraction du texte
# ----------------------------------------------------------------------------
def extract_text(data: bytes, ext: str) -> str:
    ext = ext.lower()
    if ext == ".pdf":
        from pypdf import PdfReader
        from pypdf.errors import PdfReadError

        try:
            reader = PdfReader(io.BytesIO(data))
            if reader.is_encrypted:
                raise ValueError("Ce PDF est protégé par mot de passe.")
            pages = [(page.extract_text() or "") for page in reader.pages]
        except PdfReadError as exc:
            raise ValueError("Impossible de lire ce PDF (fichier corrompu ?).") from exc
        text = "\n\n".join(pages)
        if not text.strip():
            raise ValueError("Ce PDF ne contient pas de texte extractible (document scanné ?).")
        return text
    if ext == ".docx":
        try:
            import docx
        except ImportError as exc:
            raise ValueError("La lecture des fichiers .docx nécessite le paquet python-docx.") from exc
        try:
            document = docx.Document(io.BytesIO(data))
        except Exception as exc:  # noqa: BLE001
            raise ValueError("Impossible de lire ce fichier Word.") from exc
        return "\n".join(p.text for p in document.paragraphs)
    for encoding in ("utf-8-sig", "cp1252"):
        try:
            return data.decode(encoding)
        except UnicodeDecodeError:
            continue
    return data.decode("utf-8", errors="replace")


# ----------------------------------------------------------------------------
# Découpage en passages
# ----------------------------------------------------------------------------
TARGET_CHARS = 700
MIN_CHARS = 200
HARD_MAX_CHARS = 1400


def _windows(line: str, size: int):
    if len(line) <= size:
        yield line
        return
    start = 0
    while start < len(line):
        end = min(start + size, len(line))
        if end < len(line):
            cut = line.rfind(" ", start, end)
            end = cut if cut > start + size // 2 else end
        yield line[start:end]
        start = end


def chunk_text(text: str) -> list[dict]:
    """Découpe en passages cohérents (≈700 caractères) en gardant les numéros de lignes d'origine."""
    chunks: list[dict] = []
    buf: list[tuple[int, str]] = []
    size = 0

    def flush():
        nonlocal buf, size
        lines = [(n, t) for n, t in buf if t.strip()]
        if lines:
            content = "\n".join(t for _, t in buf).strip()
            chunks.append({
                "position": len(chunks),
                "content": content,
                "start_line": lines[0][0],
                "end_line": lines[-1][0],
            })
        buf, size = [], 0

    for number, line in enumerate(text.splitlines(), start=1):
        if not line.strip():
            if buf and size >= MIN_CHARS:
                flush()
            elif buf:
                buf.append((number, ""))
            continue
        for piece in _windows(line.rstrip(), HARD_MAX_CHARS):
            if buf and size + len(piece) > HARD_MAX_CHARS:
                flush()
            buf.append((number, piece))
            size += len(piece) + 1
        if size >= TARGET_CHARS and line.rstrip().endswith((".", "!", "?", ":", ";")):
            flush()
    flush()
    return chunks


# ----------------------------------------------------------------------------
# Recherche BM25
# ----------------------------------------------------------------------------
def tokenize(text: str) -> list[str]:
    text = unicodedata.normalize("NFKD", text.lower())
    text = "".join(c for c in text if not unicodedata.combining(c))
    return [t for t in re.findall(r"[a-z0-9]{2,}", text) if t not in STOPWORDS]


def retrieve(question: str, document_ids: list[int], per_source: int = 2, max_total: int = 8) -> list[dict]:
    """Retourne les meilleurs passages des documents sélectionnés (au plus `per_source` par document)."""
    document_ids = [int(i) for i in document_ids]
    if not document_ids:
        return []
    placeholders = ",".join("?" for _ in document_ids)
    rows = db.query(
        f"""SELECT c.id AS chunk_id, c.document_id, c.position, c.content, c.start_line, c.end_line,
                   d.name AS document_name
            FROM chunks c JOIN documents d ON d.id = c.document_id
            WHERE c.document_id IN ({placeholders})
            ORDER BY c.document_id, c.position""",
        tuple(document_ids),
    )
    if not rows:
        return []

    docs_tokens = [tokenize(r["content"]) for r in rows]
    n_docs = len(rows)
    avg_len = sum(len(t) for t in docs_tokens) / n_docs or 1
    doc_freq: Counter = Counter()
    for tokens in docs_tokens:
        doc_freq.update(set(tokens))
    query_terms = set(tokenize(question))

    k1, b = 1.5, 0.75
    scored: list[tuple[float, dict]] = []
    for row, tokens in zip(rows, docs_tokens):
        freq = Counter(tokens)
        score = 0.0
        for term in query_terms:
            f = freq.get(term, 0)
            if not f:
                continue
            idf = math.log(1 + (n_docs - doc_freq[term] + 0.5) / (doc_freq[term] + 0.5))
            score += idf * f * (k1 + 1) / (f + k1 * (1 - b + b * len(tokens) / avg_len))
        scored.append((score, row))

    selected: list[tuple[float, dict]] = []
    for doc_id in document_ids:
        doc_rows = [(s, r) for s, r in scored if r["document_id"] == doc_id]
        if not doc_rows:
            continue
        best = sorted((x for x in doc_rows if x[0] > 0), key=lambda x: -x[0])[:per_source]
        if not best:  # aucun mot en commun : on propose le début du document
            best = doc_rows[:per_source]
        selected.extend(best)

    selected.sort(key=lambda x: -x[0])
    selected = selected[:max_total]
    return [
        {
            "label": f"S{i}",
            "document_id": row["document_id"],
            "document_name": row["document_name"],
            "chunk_id": row["chunk_id"],
            "position": row["position"],
            "start_line": row["start_line"],
            "end_line": row["end_line"],
            "score": round(score, 3),
            "text": row["content"],
        }
        for i, (score, row) in enumerate(selected, start=1)
    ]
